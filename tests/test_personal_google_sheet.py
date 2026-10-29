"""Tests for the personal Google-Sheet dashboard importer and exporter.

Run: uv run --with openpyxl python -m unittest tests/test_personal_google_sheet.py
"""

import importlib.util
import io
import json
import subprocess
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]


def load(name):
    spec = importlib.util.spec_from_file_location(name.replace("-", "_"), str(ROOT / "scripts" / (name + ".py")))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


PIPELINE_HEADERS = ["Company", "Role", "Pipeline status", "Current stage", "Next action / event",
                    "Follow-up date", "Waiting on"]


def current_layout(path, pipeline_headers=None):
    workbook = Workbook()
    workbook.remove(workbook.active)
    opportunities = workbook.create_sheet("Opportunities")
    opportunities.append(["Opportunity rank", "Company", "Tier", "Weighted score", "Gate", "Decision rationale"])
    opportunities.append([3, "Acme", "A", 4.2, "Conditional", "Strong fit"])
    pipeline = workbook.create_sheet("Pipeline")
    headers = pipeline_headers or (["Role", "Company"] + PIPELINE_HEADERS[2:] + ["Status category (2026-10-28)"])
    pipeline.append(headers)
    row = {"Company": "Acme", "Role": "Engineer", "Pipeline status": "On hold", "Current stage": "Screen",
           "Next action / event": "Wait", "Follow-up date": None, "Waiting on": "None",
           "Status category (2026-10-28)": "6 On hold"}
    pipeline.append([row.get(header) for header in headers])
    pipeline.append([None] * len(headers))
    workbook.create_sheet("Notes & Sources").append(["Note"])
    workbook.save(str(path))


class ImporterTests(unittest.TestCase):
    def setUp(self):
        self.importer = load("sync-google-sheet")
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_current_layout_maps_by_header_name_and_keeps_extra_columns(self):
        current_layout(self.root / "s.xlsx")
        snapshot = self.importer.sync_workbook(self.root / "s.xlsx", self.root / "out.json")
        row = snapshot["sheets"]["Pipeline"][0]
        self.assertEqual(len(snapshot["sheets"]["Pipeline"]), 1)  # blank row dropped
        self.assertEqual((row["Company"], row["Role"]), ("Acme", "Engineer"))
        self.assertEqual(row["Status category (2026-10-28)"], "6 On hold")
        self.assertEqual(snapshot["sheets"]["2026 Jobs V4"][0]["Weighted score"], 4.2)
        self.assertEqual(snapshot["sectionSources"]["2026 Jobs V4"], "Opportunities")
        self.assertEqual(snapshot["sheets"]["Interactions"], [])
        self.assertEqual(json.loads((self.root / "out.json").read_text(encoding="utf-8")), snapshot)

    def test_missing_pipeline_header_fails_without_touching_output(self):
        current_layout(self.root / "s.xlsx", pipeline_headers=["Company", "Role", "Status"])
        out = self.root / "out.json"
        out.write_text("keep me", encoding="utf-8")
        with self.assertRaises(ValueError):
            self.importer.sync_workbook(self.root / "s.xlsx", out)
        self.assertEqual(out.read_text(encoding="utf-8"), "keep me")

    def test_optional_sections_carry_forward_previous_data(self):
        current_layout(self.root / "s.xlsx")
        out = self.root / "out.json"
        out.write_text(json.dumps({"sheets": {"Diligence": [{"Company": "Acme", "Main risk": "x"}]}}), encoding="utf-8")
        snapshot = self.importer.sync_workbook(self.root / "s.xlsx", out)
        self.assertEqual(snapshot["sheets"]["Diligence"], [{"Company": "Acme", "Main risk": "x"}])
        self.assertEqual(snapshot["sectionSources"]["Diligence"], "carried-forward")


class FakeResponse(object):
    def __init__(self, body):
        self.body = body

    def read(self, _limit=-1):
        return self.body

    def __enter__(self):
        return self

    def __exit__(self, *args):
        return False


class ExporterTests(unittest.TestCase):
    def setUp(self):
        self.exporter = load("export-google-sheet")

    def test_missing_keychain_item_reports_credentials_stage(self):
        def runner(argv, **kwargs):
            raise subprocess.CalledProcessError(44, argv)
        with self.assertRaises(self.exporter.ExportError) as caught:
            self.exporter.export_and_sync(runner=runner, urlopen=lambda *a, **k: None, sheet_id="sheet-123", adc=lambda: None)
        self.assertEqual(caught.exception.stage, "credentials")

    def test_happy_path_uses_py36_safe_subprocess_kwargs_and_cleans_up(self):
        calls = []
        credential = json.dumps({"client_id": "id", "client_secret": "s", "refresh_token": "r"})

        def runner(argv, **kwargs):
            calls.append((argv, kwargs))
            self.assertNotIn("capture_output", kwargs)
            self.assertNotIn("text", kwargs)
            self.assertEqual(kwargs["stdout"], subprocess.PIPE)
            return subprocess.CompletedProcess(argv, 0, stdout=credential, stderr="")

        def urlopen(request, timeout):
            if "oauth2" in request.full_url:
                return FakeResponse(b'{"access_token": "tok"}')
            self.assertIn("/files/sheet-123/export", request.full_url)
            self.assertEqual(request.get_header("Authorization"), "Bearer tok")
            return FakeResponse(b"PK\x03\x04workbook")

        with tempfile.TemporaryDirectory() as tmp:
            self.exporter.export_and_sync(output_path=Path(tmp) / "o.json", runner=runner, urlopen=urlopen, sheet_id="sheet-123", adc=lambda: None)
        self.assertEqual(calls[0][0], self.exporter.KEYCHAIN_FIND_ARGV)
        self.assertEqual(Path(calls[1][0][1]), self.exporter.IMPORTER_SCRIPT)
        self.assertFalse(Path(calls[1][0][2]).exists())

    def test_sheet_id_comes_from_env_or_ignored_file_never_the_repo(self):
        self.assertEqual(self.exporter.spreadsheet_id({"CAREER_OPS_SHEET_ID": " abc_D-1 "}), "abc_D-1")
        with tempfile.TemporaryDirectory() as tmp:
            id_file = Path(tmp) / "id.txt"
            with self.assertRaises(self.exporter.ExportError) as caught:
                self.exporter.spreadsheet_id({}, id_file)
            self.assertEqual(caught.exception.stage, "config")
            id_file.write_text("fromfile\n")
            self.assertEqual(self.exporter.spreadsheet_id({}, id_file), "fromfile")
            with self.assertRaises(self.exporter.ExportError):
                self.exporter.spreadsheet_id({"CAREER_OPS_SHEET_ID": "../evil?x"}, id_file)
        self.assertFalse(hasattr(self.exporter, "SPREADSHEET_ID"))

    def test_adc_is_preferred_and_sends_quota_project(self):
        seen = {}

        def urlopen(request, timeout):
            seen["auth"] = request.get_header("Authorization")
            seen["quota"] = request.get_header("X-goog-user-project")
            return FakeResponse(b"PK\x03\x04workbook")

        def runner(argv, **kwargs):
            if argv[0] == "security":
                raise AssertionError("Keychain must not be read when ADC works")
            return subprocess.CompletedProcess(argv, 0, stdout="", stderr="")

        with tempfile.TemporaryDirectory() as tmp:
            self.exporter.export_and_sync(output_path=Path(tmp) / "o.json", runner=runner, urlopen=urlopen,
                                          sheet_id="sheet-123", adc=lambda: {"token": "adc-tok", "quota_project": "my-proj"})
        self.assertEqual(seen, {"auth": "Bearer adc-tok", "quota": "my-proj"})

    def test_adc_helper_handles_missing_and_scopes(self):
        class DefaultCredentialsError(Exception):
            pass

        def missing(scopes):
            raise DefaultCredentialsError("none")
        self.assertIsNone(self.exporter.adc_token(default=missing, request_factory=object))

        class Creds:
            token = "t"
            quota_project_id = None
            def refresh(self, request):
                pass
        captured = {}
        def ok(scopes):
            captured["scopes"] = scopes
            return Creds(), None
        self.assertEqual(self.exporter.adc_token(default=ok, request_factory=object)["token"], "t")
        self.assertIn("https://www.googleapis.com/auth/drive.readonly", captured["scopes"])
        self.assertIn("https://www.googleapis.com/auth/spreadsheets.readonly", captured["scopes"])
        self.assertEqual(self.exporter.classify_http_error(403, b'{"reason": "ACCESS_TOKEN_SCOPE_INSUFFICIENT"}').stage, "scopes")
        self.assertEqual(self.exporter.classify_http_error(403, b'Drive API has not been used in project 1').stage, "api-disabled")
        self.assertEqual(self.exporter.classify_http_error(404, b"").stage, "not-found")

    def test_main_prints_secret_free_json_stage(self):
        def boom(**kwargs):
            raise self.exporter.ExportError("download", "Unable to export")
        original = self.exporter.export_and_sync
        self.exporter.export_and_sync = boom
        stderr = io.StringIO()
        import contextlib
        try:
            with contextlib.redirect_stderr(stderr):
                self.assertEqual(self.exporter.main(), 1)
        finally:
            self.exporter.export_and_sync = original
        self.assertEqual(json.loads(stderr.getvalue())["stage"], "download")


if __name__ == "__main__":
    unittest.main()
