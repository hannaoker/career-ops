#!/usr/bin/env python3
# /// script
# requires-python = ">=3.9"
# dependencies = [
#     "openpyxl>=3.0",
# ]
# ///
"""Convert a local Google Sheets XLSX export into a private JSON snapshot.

The snapshot always exposes four logical sections (``Pipeline``,
``Interactions``, ``2026 Jobs V4`` and ``Diligence``) so the dashboard
contract stays stable while the workbook layout evolves:

* ``Pipeline`` is required and comes from the ``Pipeline`` tab.
* ``2026 Jobs V4`` (rankings) comes from the first tab found among
  ``Opportunities`` (current layout) and ``2026 Jobs V4`` (legacy layout).
* ``Interactions`` and ``Diligence`` are optional. When the workbook has no
  usable source tab, the previous snapshot's section is carried forward, or an
  empty list is used, instead of failing the whole sync.

Columns are always mapped by header name, so extra or reordered columns (for
example ``Status category (2026-10-28)``) are imported as-is.
"""

import argparse
import json
import os
import tempfile
from datetime import date, datetime, time, timezone
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUTPUT = PROJECT_ROOT / "data" / "personal-dashboard.json"
SNAPSHOT_SECTIONS = ("Pipeline", "Interactions", "2026 Jobs V4", "Diligence")
REQUIRED_SECTIONS = ("Pipeline",)
# Snapshot section -> workbook tab names to try, in priority order.
SECTION_SOURCE_TABS = {
    "Pipeline": ("Pipeline",),
    "Interactions": ("Interactions",),
    "2026 Jobs V4": ("Opportunities", "2026 Jobs V4"),
    "Diligence": ("Diligence",),
}
REQUIRED_HEADERS = {
    "Pipeline": (
        "Company",
        "Role",
        "Pipeline status",
        "Current stage",
        "Next action / event",
        "Follow-up date",
        "Waiting on",
    ),
    "Interactions": ("Date", "Company", "Interaction type", "Summary"),
    "2026 Jobs V4": ("Company",),
    "Diligence": ("Company", "Main risk", "Financial safety score (1–5)"),
}


def _is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def _json_value(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _cell_value(formula_cell, cached_cell):
    if formula_cell.data_type == "f":
        return _json_value(cached_cell.value)
    return _json_value(formula_cell.value)


def _sheet_rows(formula_sheet, cached_sheet, required_headers):
    formula_rows = formula_sheet.iter_rows()
    cached_rows = cached_sheet.iter_rows()
    try:
        formula_header = next(formula_rows)
        cached_header = next(cached_rows)
    except StopIteration:
        raise ValueError(
            f"Sheet {formula_sheet.title!r} is missing required header(s): "
            f"{', '.join(required_headers)}"
        )

    raw_headers = [
        _cell_value(formula_cell, cached_cell)
        for formula_cell, cached_cell in zip(formula_header, cached_header)
    ]
    while raw_headers and _is_blank(raw_headers[-1]):
        raw_headers.pop()
    if not raw_headers:
        raise ValueError(
            f"Sheet {formula_sheet.title!r} is missing required header(s): "
            f"{', '.join(required_headers)}"
        )
    if any(_is_blank(header) for header in raw_headers):
        raise ValueError(f"Sheet {formula_sheet.title!r} has a blank header")

    headers = [str(header).strip() for header in raw_headers]
    if len(headers) != len(set(headers)):
        raise ValueError(f"Sheet {formula_sheet.title!r} has duplicate headers")
    missing_headers = [header for header in required_headers if header not in headers]
    if missing_headers:
        raise ValueError(
            f"Sheet {formula_sheet.title!r} is missing required header(s): "
            f"{', '.join(missing_headers)}"
        )

    rows = []
    for formula_row, cached_row in zip(formula_rows, cached_rows):
        values = [
            _cell_value(formula_cell, cached_cell)
            for formula_cell, cached_cell in zip(formula_row[: len(headers)], cached_row[: len(headers)])
        ]
        values.extend([None] * (len(headers) - len(values)))
        if all(_is_blank(value) for value in values):
            continue
        rows.append(dict(zip(headers, values)))
    return rows


def _previous_sections(output_path):
    """Best-effort read of the existing snapshot so optional sections survive."""
    try:
        previous = json.loads(output_path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}
    sheets = previous.get("sheets") if isinstance(previous, dict) else None
    if not isinstance(sheets, dict):
        return {}
    return {
        name: rows
        for name, rows in sheets.items()
        if isinstance(rows, list) and all(isinstance(row, dict) for row in rows)
    }


def _write_atomically(snapshot, output_path):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=output_path.parent,
            prefix=f".{output_path.name}.",
            suffix=".tmp",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
            json.dump(snapshot, temporary, ensure_ascii=False, indent=2, allow_nan=False)
            temporary.write("\n")
            temporary.flush()
            os.fsync(temporary.fileno())
        os.replace(temporary_path, output_path)
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def sync_workbook(workbook_path, output_path=DEFAULT_OUTPUT):
    workbook_path = Path(workbook_path).expanduser().resolve()
    output_path = Path(output_path).expanduser().resolve()

    formula_workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    cached_workbook = load_workbook(workbook_path, data_only=True, read_only=False)
    try:
        tab_names = set(formula_workbook.sheetnames)
        missing = [
            name
            for name in REQUIRED_SECTIONS
            if not any(tab in tab_names for tab in SECTION_SOURCE_TABS[name])
        ]
        if missing:
            raise ValueError(f"Workbook is missing required tab(s): {', '.join(missing)}")

        source = {"type": "xlsx", "filename": workbook_path.name}
        title = formula_workbook.properties.title
        if title and title.strip():
            source["title"] = title.strip()

        previous = None
        sheets = {}
        section_sources = {}
        warnings = []
        for name in SNAPSHOT_SECTIONS:
            tab = next((tab for tab in SECTION_SOURCE_TABS[name] if tab in tab_names), None)
            if tab is not None:
                try:
                    sheets[name] = _sheet_rows(
                        formula_workbook[tab],
                        cached_workbook[tab],
                        REQUIRED_HEADERS[name],
                    )
                    section_sources[name] = tab
                    continue
                except ValueError as exc:
                    if name in REQUIRED_SECTIONS:
                        raise
                    warnings.append(str(exc))
            # Optional section without a usable source tab: keep prior data.
            if previous is None:
                previous = _previous_sections(output_path)
            if name in previous:
                sheets[name] = previous[name]
                section_sources[name] = "carried-forward"
            else:
                sheets[name] = []
                section_sources[name] = "missing"

        snapshot = {
            "schemaVersion": 1,
            "generatedAt": datetime.now(timezone.utc).isoformat().replace("+00:00", "Z"),
            "source": source,
            "sectionSources": section_sources,
            "sheets": sheets,
        }
        if warnings:
            snapshot["warnings"] = warnings
        _write_atomically(snapshot, output_path)
        return snapshot
    finally:
        formula_workbook.close()
        cached_workbook.close()


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="Path to an authenticated .xlsx export")
    parser.add_argument(
        "output",
        nargs="?",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"Private JSON snapshot path (default: {DEFAULT_OUTPUT})",
    )
    args = parser.parse_args(argv)
    sync_workbook(args.workbook, args.output)


if __name__ == "__main__":
    main()
